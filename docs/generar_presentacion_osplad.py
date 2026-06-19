#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera la presentación (PPTX) explicativa del flujo OSPLAD para las farmacias.
Uso:  /tmp/pptx_venv/bin/python docs/generar_presentacion_osplad.py
Salida: docs/Instructivo_Farmacias_OSPLAD.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Paleta (coherente con el módulo: violeta/azul)
MORADO  = RGBColor(0x66, 0x7E, 0xEA)
MORADO2 = RGBColor(0x76, 0x4B, 0xA2)
GRIS    = RGBColor(0x44, 0x44, 0x44)
GRIS_CL = RGBColor(0x6B, 0x6B, 0x6B)
BLANCO  = RGBColor(0xFF, 0xFF, 0xFF)
AMARILLO= RGBColor(0xF0, 0xAD, 0x4E)
AZUL    = RGBColor(0x31, 0x9A, 0xD4)
VERDE   = RGBColor(0x28, 0xA7, 0x45)
FONDO   = RGBColor(0xF5, 0xF5, 0xF7)

prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def fondo(slide, color=BLANCO):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def caja(slide, x, y, w, h, color=None, line=None, radius=False):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    if color is not None:
        shp.fill.solid(); shp.fill.fore_color.rgb = color
    else:
        shp.fill.background()
    if line is not None:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def texto(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6):
    """runs: lista de párrafos; cada párrafo es lista de (texto, size, bold, color)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, par in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        for (t, sz, bold, col) in par:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = bold; r.font.color.rgb = col
            r.font.name = "Calibri"
    return tb


def banda(slide, color=MORADO, h=Inches(1.15)):
    caja(slide, 0, 0, SW, h, color=color)


# ---------------------------------------------------------------- Slide 1: portada
s = prs.slides.add_slide(BLANK); fondo(s, MORADO)
caja(s, 0, Inches(2.9), SW, Inches(1.7), color=MORADO2)
texto(s, Inches(1), Inches(2.95), Inches(11.3), Inches(1.6),
      [[("Sistema OSPLAD", 44, True, BLANCO)],
       [("Instructivo para Farmacias — Gestión de pedidos", 22, False, RGBColor(0xEA,0xE6,0xF7))]],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
texto(s, Inches(1), Inches(6.3), Inches(11.3), Inches(0.7),
      [[("Global Médica S.A.  ·  Sistema SISCON", 16, False, RGBColor(0xD9,0xD2,0xF0))]],
      align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- Slide 2: qué es
s = prs.slides.add_slide(BLANK); fondo(s, FONDO); banda(s)
texto(s, Inches(0.6), Inches(0.18), Inches(12), Inches(0.8),
      [[("¿Qué vas a hacer en el sistema?", 30, True, BLANCO)]], anchor=MSO_ANCHOR.MIDDLE)
texto(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.2),
      [[("El sistema te muestra los pedidos de medicación de los afiliados de OSPLAD "
         "asignados a tu farmacia. Vos solo tenés que confirmar la entrega cuando el "
         "afiliado retira. El resto del proceso es automático.", 20, False, GRIS)]])
puntos = [
    ("Ingresás con tu usuario y contraseña.", "Solo ves los pedidos de tu farmacia."),
    ("El sistema avisa al afiliado por WhatsApp.", "Cuando el pedido está en camino y cuando se entrega."),
    ("Cuando el afiliado retira, marcás 'Entregar'.", "Se genera el comprobante (PDF) para firmar."),
]
y = Inches(3.1)
for i, (a, b) in enumerate(puntos):
    caja(s, Inches(0.8), y, Inches(11.7), Inches(1.05), color=BLANCO, radius=True)
    caja(s, Inches(0.8), y, Inches(0.16), Inches(1.05), color=MORADO, radius=False)
    texto(s, Inches(1.15), y+Inches(0.12), Inches(11), Inches(0.85),
          [[("%d. %s" % (i+1, a), 18, True, MORADO2)],
           [(b, 15, False, GRIS_CL)]])
    y += Inches(1.2)

# ---------------------------------------------------------------- Slide 3: las 3 etapas
s = prs.slides.add_slide(BLANK); fondo(s, FONDO); banda(s)
texto(s, Inches(0.6), Inches(0.18), Inches(12), Inches(0.8),
      [[("El flujo: 3 etapas", 30, True, BLANCO)]], anchor=MSO_ANCHOR.MIDDLE)
etapas = [
    ("PENDIENTES", AMARILLO, "El pedido fue habilitado.",
     "Lo podés VER e IMPRIMIR. No hay que hacer nada todavía."),
    ("EN TRÁNSITO", AZUL, "El pedido está siendo preparado / en camino.",
     "El afiliado ya recibió un WhatsApp. Acá vas a marcar la entrega."),
    ("ENTREGAS", VERDE, "El pedido fue entregado al afiliado.",
     "Queda registrado con la fecha. El afiliado recibe el aviso de entrega."),
]
x = Inches(0.7); w = Inches(3.95)
for (titulo, col, sub, desc) in etapas:
    caja(s, x, Inches(1.6), w, Inches(4.7), color=BLANCO, radius=True)
    caja(s, x, Inches(1.6), w, Inches(0.95), color=col, radius=True)
    texto(s, x, Inches(1.6), w, Inches(0.95),
          [[(titulo, 20, True, BLANCO)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    texto(s, x+Inches(0.25), Inches(2.8), w-Inches(0.5), Inches(3.3),
          [[(sub, 16, True, GRIS)], [("", 6, False, GRIS)], [(desc, 15, False, GRIS_CL)]])
    x += w + Inches(0.3)
# flechas
for ax in (Inches(4.72), Inches(8.97)):
    texto(s, ax, Inches(3.4), Inches(0.3), Inches(0.8),
          [[("→", 30, True, MORADO2)]], align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- Slide 4: ingreso
s = prs.slides.add_slide(BLANK); fondo(s, FONDO); banda(s)
texto(s, Inches(0.6), Inches(0.18), Inches(12), Inches(0.8),
      [[("1) Ingreso al sistema", 30, True, BLANCO)]], anchor=MSO_ANCHOR.MIDDLE)
texto(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(3),
      [[("Entrá a la dirección del sistema y poné tu correo y contraseña.", 20, False, GRIS)],
       [("", 8, False, GRIS)],
       [("• Usuario: ", 18, True, MORADO2), ("el correo de tu farmacia.", 18, False, GRIS)],
       [("• Contraseña: ", 18, True, MORADO2), ("la que te entregamos (podés cambiarla luego).", 18, False, GRIS)],
       [("", 8, False, GRIS)],
       [("Al ingresar verás directamente tu módulo OSPLAD, con las pestañas "
         "PENDIENTES, En tránsito y Entregas. Solo aparecen los pedidos de tu farmacia.", 18, False, GRIS)]])
caja(s, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.2), color=RGBColor(0xFF,0xF3,0xCD), radius=True)
texto(s, Inches(1.1), Inches(5.5), Inches(11.1), Inches(1),
      [[("Importante: ", 16, True, RGBColor(0x85,0x64,0x04)),
        ("no necesitás cargar nada. El sistema ya trae los pedidos habilitados de OSPLAD.",
         16, False, RGBColor(0x85,0x64,0x04))]], anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- Slide 5: entregar
s = prs.slides.add_slide(BLANK); fondo(s, FONDO); banda(s)
texto(s, Inches(0.6), Inches(0.18), Inches(12), Inches(0.8),
      [[("2) Entregar un pedido", 30, True, BLANCO)]], anchor=MSO_ANCHOR.MIDDLE)
pasos = [
    "Andá a la pestaña 'En tránsito'.",
    "Buscá al afiliado que viene a retirar (podés filtrar por nombre o DNI).",
    "Tildá el/los pedidos. Si son varios del mismo afiliado y mismo remito, se entregan juntos.",
    "Hacé clic en 'Entregar seleccionados' y confirmá.",
    "Se abre el comprobante en PDF para imprimir y que el afiliado firme.",
]
y = Inches(1.55)
for i, p in enumerate(pasos):
    caja(s, Inches(0.8), y, Inches(0.55), Inches(0.55), color=MORADO, radius=True)
    texto(s, Inches(0.8), y, Inches(0.55), Inches(0.55),
          [[(str(i+1), 18, True, BLANCO)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    texto(s, Inches(1.55), y, Inches(10.9), Inches(0.6),
          [[(p, 18, False, GRIS)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.92)
caja(s, Inches(0.8), Inches(6.25), Inches(11.7), Inches(0.95), color=RGBColor(0xD4,0xED,0xDA), radius=True)
texto(s, Inches(1.1), Inches(6.3), Inches(11.1), Inches(0.85),
      [[("Entrega unificada: ", 16, True, RGBColor(0x15,0x57,0x24)),
        ("un mismo remito puede agrupar varios pedidos del mismo afiliado.",
         16, False, RGBColor(0x15,0x57,0x24))]], anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- Slide 6: WhatsApp
s = prs.slides.add_slide(BLANK); fondo(s, FONDO); banda(s)
texto(s, Inches(0.6), Inches(0.18), Inches(12), Inches(0.8),
      [[("Avisos automáticos al afiliado (WhatsApp)", 28, True, BLANCO)]], anchor=MSO_ANCHOR.MIDDLE)
texto(s, Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.7),
      [[("El sistema envía los mensajes solo. La farmacia no escribe nada.", 18, False, GRIS)]])
msgs = [
    ("Cuando el pedido está en proceso / camino",
     "Buenos días, nos contactamos desde Global Médica. Queremos informarle que su "
     "medicación está siendo procesada. Podrá retirarla por la sucursal en 96 hs (4 días hábiles)."),
    ("Cuando marcás la entrega",
     "Buenos días, nos contactamos desde Global Médica. Queremos informarle que su pedido "
     "[N°] ha sido entregado el día [fecha] en la farmacia [nombre]. ¡Muchas gracias!"),
]
y = Inches(2.15)
for (tit, cuerpo) in msgs:
    caja(s, Inches(0.8), y, Inches(11.7), Inches(2.05), color=BLANCO, radius=True)
    caja(s, Inches(0.8), y, Inches(0.16), Inches(2.05), color=VERDE)
    texto(s, Inches(1.15), y+Inches(0.15), Inches(11.1), Inches(1.8),
          [[(tit, 17, True, VERDE)],
           [("“%s”" % cuerpo, 15, False, GRIS)]])
    y += Inches(2.2)

# ---------------------------------------------------------------- Slide 7: comprobante + cierre
s = prs.slides.add_slide(BLANK); fondo(s, FONDO); banda(s)
texto(s, Inches(0.6), Inches(0.18), Inches(12), Inches(0.8),
      [[("Comprobante de entrega (PDF)", 28, True, BLANCO)]], anchor=MSO_ANCHOR.MIDDLE)
texto(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(3),
      [[("Al confirmar la entrega se genera un comprobante en PDF, listo para imprimir:", 19, False, GRIS)],
       [("", 8, False, GRIS)],
       [("• Nombre de la obra social (OSPLAD) y datos del afiliado.", 17, False, GRIS)],
       [("• Detalle de la medicación entregada y remito.", 17, False, GRIS)],
       [("• Fecha de entrega y espacio para la firma del afiliado.", 17, False, GRIS)],
       [("", 8, False, GRIS)],
       [("Imprimilo, hacelo firmar y guardalo como constancia de la entrega.", 17, True, MORADO2)]])
caja(s, 0, Inches(6.55), SW, Inches(0.95), color=MORADO2)
texto(s, Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.95),
      [[("¿Dudas? Comunicate con Global Médica.  ·  ¡Gracias por trabajar con nosotros!",
         16, True, BLANCO)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Instructivo_Farmacias_OSPLAD.pptx")
prs.save(out)
print("OK:", out, "(%d diapositivas)" % len(prs.slides._sldIdLst))
